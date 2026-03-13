"""
All-Documents Reference Database Indexer.

Reads multiple Excel files from the All_Docs SharePoint folder, builds an
in-memory lookup index keyed by Bates number, and provides fast metadata
retrieval. The all-docs files are READ-ONLY reference material.
"""

import logging
import re
from pathlib import Path
from typing import Optional

import openpyxl

from config.settings import ALL_DOCS_COLUMNS, ALL_DOCS_FOLDER, AUTO_POPULATE_COLUMNS, BATES_PATTERN
from graph_api_client import GraphAPIClient

logger = logging.getLogger(__name__)


class AllDocsIndexer:
    """Index ~300k documents across multiple Excel files for fast Bates lookup."""

    def __init__(self, client: GraphAPIClient):
        self.client = client
        # Main index: {bates_number: {col: value, ...}}
        self._index: dict[str, dict] = {}
        self._loaded = False

    @property
    def count(self) -> int:
        return len(self._index)

    def load(self) -> None:
        """Download and index all Excel files from the All_Docs folder."""
        logger.info("Loading all-docs index from SharePoint...")
        items = self.client.list_folder(ALL_DOCS_FOLDER)
        xlsx_files = [
            item["name"] for item in items
            if "file" in item and item["name"].lower().endswith(".xlsx")
        ]
        logger.info(f"Found {len(xlsx_files)} all-docs Excel files.")

        for filename in xlsx_files:
            self._load_single_file(filename)

        self._loaded = True
        logger.info(f"All-docs index loaded: {self.count} documents indexed.")

    def _load_single_file(self, filename: str) -> None:
        """Load a single all-docs Excel file into the index."""
        relative_path = f"{ALL_DOCS_FOLDER}/{filename}"
        logger.info(f"  Loading {filename}...")

        try:
            stream = self.client.download_file_to_stream(relative_path)
            wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        except Exception as e:
            logger.error(f"  Failed to load {filename}: {e}")
            return

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            # Find header row — look for "Bates" column
            header_row = rows[0]
            headers = [str(h).strip() if h else "" for h in header_row]

            if "Bates" not in headers:
                logger.warning(f"  Sheet '{sheet_name}' in {filename} has no 'Bates' column, skipping.")
                continue

            bates_col_idx = headers.index("Bates")

            # Map column indices to column names
            col_map = {}
            for i, h in enumerate(headers):
                if h in ALL_DOCS_COLUMNS:
                    col_map[i] = h

            # Index each row
            for row in rows[1:]:
                if not row or len(row) <= bates_col_idx:
                    continue
                bates = row[bates_col_idx]
                if not bates:
                    continue
                bates = str(bates).strip()
                if not bates:
                    continue

                record = {}
                for i, col_name in col_map.items():
                    if i < len(row):
                        val = row[i]
                        record[col_name] = val if val is not None else ""
                    else:
                        record[col_name] = ""

                self._index[bates] = record

        wb.close()

    def load_from_local_files(self, folder_path: str) -> None:
        """
        Load all-docs index from local Excel files (for offline/testing use).

        Args:
            folder_path: Path to a local folder containing all-docs .xlsx files.
        """
        logger.info(f"Loading all-docs index from local folder: {folder_path}")
        folder = Path(folder_path)
        xlsx_files = sorted(folder.glob("*.xlsx"))
        logger.info(f"Found {len(xlsx_files)} files.")

        for filepath in xlsx_files:
            try:
                wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            except Exception as e:
                logger.error(f"  Failed to load {filepath.name}: {e}")
                continue

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                headers = [str(h).strip() if h else "" for h in rows[0]]
                if "Bates" not in headers:
                    continue

                bates_col_idx = headers.index("Bates")
                col_map = {i: h for i, h in enumerate(headers) if h in ALL_DOCS_COLUMNS}

                for row in rows[1:]:
                    if not row or len(row) <= bates_col_idx:
                        continue
                    bates = row[bates_col_idx]
                    if not bates:
                        continue
                    bates = str(bates).strip()
                    if not bates:
                        continue

                    record = {}
                    for i, col_name in col_map.items():
                        if i < len(row):
                            val = row[i]
                            record[col_name] = val if val is not None else ""
                        else:
                            record[col_name] = ""
                    self._index[bates] = record

            wb.close()

        self._loaded = True
        logger.info(f"All-docs index loaded: {self.count} documents indexed.")

    def lookup(self, bates_number: str) -> Optional[dict]:
        """
        Look up a Bates number and return its metadata.

        Returns None if not found.
        """
        if not self._loaded:
            raise RuntimeError("Index not loaded. Call load() or load_from_local_files() first.")
        return self._index.get(bates_number.strip())

    def get_auto_populate_data(self, bates_number: str) -> Optional[dict]:
        """
        Look up a Bates number and return only the columns that should be
        auto-populated into the exhibit list.
        """
        record = self.lookup(bates_number)
        if not record:
            return None
        return {col: record.get(col, "") for col in AUTO_POPULATE_COLUMNS}

    def search(self, query: str) -> list[tuple[str, dict]]:
        """
        Search the index by partial Bates number or metadata value.

        Returns list of (bates_number, record) tuples.
        """
        if not self._loaded:
            raise RuntimeError("Index not loaded.")

        query_lower = query.lower()
        results = []
        for bates, record in self._index.items():
            if query_lower in bates.lower():
                results.append((bates, record))
                continue
            for val in record.values():
                if val and query_lower in str(val).lower():
                    results.append((bates, record))
                    break
        return results

    def get_all_bates_numbers(self) -> set[str]:
        """Return the set of all Bates numbers in the index."""
        return set(self._index.keys())
