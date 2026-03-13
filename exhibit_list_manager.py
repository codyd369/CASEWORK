"""
Exhibit List Manager.

Manages the exhibit list Excel workbook in SharePoint:
- Reads the exhibit list and identifies rows needing metadata
- Auto-populates metadata from the all-docs index
- Checks for missing documents and highlights rows RED
- Generates document open links
- Adds new rows when new documents appear in the folder
"""

import io
import logging
import re
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import PatternFill

from config.settings import (
    AUTO_POPULATE_COLUMNS,
    BATES_PATTERN,
    DOCUMENTS_FOLDER,
    EXHIBIT_LIST_FILENAME,
)
from all_docs_indexer import AllDocsIndexer
from graph_api_client import GraphAPIClient

logger = logging.getLogger(__name__)

# Red fill for missing documents
RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
NO_FILL = PatternFill(fill_type=None)


class ExhibitListManager:
    """Manages the exhibit list workbook."""

    def __init__(self, client: GraphAPIClient, indexer: AllDocsIndexer):
        self.client = client
        self.indexer = indexer
        self._wb: Optional[openpyxl.Workbook] = None
        self._ws = None  # Active worksheet with exhibit data
        self._headers: list[str] = []
        self._header_map: dict[str, int] = {}  # column name → column index (1-based)

    def load(self) -> None:
        """Download and open the exhibit list workbook."""
        logger.info("Loading exhibit list from SharePoint...")
        stream = self.client.download_file_to_stream(EXHIBIT_LIST_FILENAME)
        self._wb = openpyxl.load_workbook(stream)
        self._find_data_sheet()
        logger.info(f"Exhibit list loaded: sheet='{self._ws.title}', "
                     f"{self._ws.max_row - 1} rows, {len(self._headers)} columns.")

    def _find_data_sheet(self) -> None:
        """
        Find the worksheet that contains the exhibit data.
        Looks for a sheet with a 'Bates' column header in row 1.
        """
        for sheet_name in self._wb.sheetnames:
            ws = self._wb[sheet_name]
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=1, column=col).value
                if val and str(val).strip() == "Bates":
                    self._ws = ws
                    self._read_headers()
                    return

        # Fallback: use the first sheet
        self._ws = self._wb.active
        self._read_headers()
        logger.warning("No sheet with 'Bates' header found; using active sheet.")

    def _read_headers(self) -> None:
        """Read column headers from row 1."""
        self._headers = []
        self._header_map = {}
        for col in range(1, self._ws.max_column + 1):
            val = self._ws.cell(row=1, column=col).value
            header = str(val).strip() if val else ""
            self._headers.append(header)
            if header:
                self._header_map[header] = col

    def _col(self, name: str) -> Optional[int]:
        """Get 1-based column index for a header name."""
        return self._header_map.get(name)

    def _get_cell_value(self, row: int, col_name: str) -> Optional[str]:
        """Get a cell's value by row number and column name."""
        col_idx = self._col(col_name)
        if not col_idx:
            return None
        val = self._ws.cell(row=row, column=col_idx).value
        return str(val).strip() if val else None

    def _set_cell_value(self, row: int, col_name: str, value) -> None:
        """Set a cell's value by row number and column name."""
        col_idx = self._col(col_name)
        if col_idx:
            self._ws.cell(row=row, column=col_idx).value = value

    # ── Core Operations ──

    def get_all_bates_numbers(self) -> list[tuple[int, str]]:
        """
        Get all Bates numbers from the exhibit list.

        Returns: list of (row_number, bates_number) tuples.
        """
        bates_col = self._col("Bates")
        if not bates_col:
            logger.error("No 'Bates' column found in exhibit list.")
            return []

        result = []
        for row in range(2, self._ws.max_row + 1):
            val = self._ws.cell(row=row, column=bates_col).value
            if val:
                result.append((row, str(val).strip()))
        return result

    def find_rows_needing_metadata(self) -> list[tuple[int, str]]:
        """
        Find rows that have a Bates number but are missing metadata.

        A row "needs metadata" if it has a Bates number and any of the
        auto-populate columns are empty.

        Returns: list of (row_number, bates_number) tuples.
        """
        rows_needing = []
        for row, bates in self.get_all_bates_numbers():
            for col_name in AUTO_POPULATE_COLUMNS:
                val = self._get_cell_value(row, col_name)
                if not val:
                    rows_needing.append((row, bates))
                    break
        return rows_needing

    def auto_populate_metadata(self) -> int:
        """
        Auto-populate metadata for rows that have a Bates number but missing data.

        Returns the number of rows updated.
        """
        rows = self.find_rows_needing_metadata()
        updated = 0

        for row, bates in rows:
            data = self.indexer.get_auto_populate_data(bates)
            if not data:
                logger.warning(f"  Row {row}: Bates '{bates}' not found in all-docs index.")
                continue

            for col_name, value in data.items():
                if value and not self._get_cell_value(row, col_name):
                    self._set_cell_value(row, col_name, value)

            updated += 1
            logger.info(f"  Row {row}: Populated metadata for '{bates}'.")

        return updated

    def check_missing_documents(self) -> list[tuple[int, str]]:
        """
        Check for exhibits whose documents are missing from the SharePoint folder.

        Highlights missing rows RED. Clears highlighting for found documents.

        Returns: list of (row_number, bates_number) for missing documents.
        """
        logger.info("Checking for missing documents in SharePoint folder...")
        # Get all files in the document folder, keyed by stem (Bates number)
        folder_files = self.client.list_files_in_folder(DOCUMENTS_FOLDER)
        folder_stems = set(folder_files.keys())

        missing = []
        bates_col = self._col("Bates")
        if not bates_col:
            return missing

        for row, bates in self.get_all_bates_numbers():
            if bates in folder_stems:
                # Document found — clear any red highlighting
                for col in range(1, self._ws.max_column + 1):
                    cell = self._ws.cell(row=row, column=col)
                    if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb == "00FF0000":
                        cell.fill = NO_FILL
            else:
                # Document missing — highlight row RED
                missing.append((row, bates))
                for col in range(1, self._ws.max_column + 1):
                    self._ws.cell(row=row, column=col).fill = RED_FILL

        logger.info(f"Missing documents: {len(missing)} out of {len(self.get_all_bates_numbers())} exhibits.")
        return missing

    def add_new_exhibits_from_folder(self) -> int:
        """
        Scan the SharePoint document folder for files not in the exhibit list.
        Add new rows for any new documents found.

        Returns the number of new exhibits added.
        """
        logger.info("Scanning document folder for new exhibits...")
        folder_files = self.client.list_files_in_folder(DOCUMENTS_FOLDER)
        existing_bates = {bates for _, bates in self.get_all_bates_numbers()}

        added = 0
        for stem, file_info in folder_files.items():
            # Check if stem looks like a Bates number
            if not re.match(BATES_PATTERN, stem):
                logger.debug(f"  Skipping non-Bates file: {file_info['name']}")
                continue

            if stem in existing_bates:
                continue

            # New document — add a row
            new_row = self._ws.max_row + 1
            self._set_cell_value(new_row, "Bates", stem)

            # Try to populate metadata from all-docs
            data = self.indexer.get_auto_populate_data(stem)
            if data:
                for col_name, value in data.items():
                    if value:
                        self._set_cell_value(new_row, col_name, value)

            added += 1
            logger.info(f"  Added new exhibit: '{stem}' at row {new_row}.")

        return added

    def add_document_links(self) -> int:
        """
        Add/update clickable document open links for each exhibit row.

        Adds a 'Document Link' column if it doesn't exist, with hyperlinks
        that open the document via SharePoint (triggering native app).
        """
        link_col_name = "Document Link"
        if link_col_name not in self._header_map:
            # Add the column
            new_col = self._ws.max_column + 1
            self._ws.cell(row=1, column=new_col).value = link_col_name
            self._header_map[link_col_name] = new_col
            self._headers.append(link_col_name)

        link_col = self._col(link_col_name)
        folder_files = self.client.list_files_in_folder(DOCUMENTS_FOLDER)
        updated = 0

        for row, bates in self.get_all_bates_numbers():
            if bates in folder_files:
                file_info = folder_files[bates]
                url = file_info["webUrl"]
                # For Office docs, append ?web=0 to force desktop app
                ext = file_info["extension"]
                if ext in (".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"):
                    url = f"{url}?web=0"

                cell = self._ws.cell(row=row, column=link_col)
                cell.value = "Open"
                cell.hyperlink = url
                cell.style = "Hyperlink"
                updated += 1

        return updated

    def save(self) -> None:
        """Save the workbook back to SharePoint."""
        logger.info("Saving exhibit list to SharePoint...")
        buffer = io.BytesIO()
        self._wb.save(buffer)
        buffer.seek(0)
        self.client.upload_file(EXHIBIT_LIST_FILENAME, buffer.read())
        logger.info("Exhibit list saved.")

    def save_local(self, path: str) -> None:
        """Save the workbook to a local file (for testing)."""
        self._wb.save(path)
        logger.info(f"Exhibit list saved locally to {path}.")

    def get_exhibit_metadata(self, bates_number: str) -> Optional[dict]:
        """Get all metadata for a specific exhibit from the exhibit list."""
        for row, bates in self.get_all_bates_numbers():
            if bates == bates_number:
                record = {}
                for header in self._headers:
                    if header:
                        record[header] = self._get_cell_value(row, header) or ""
                return record
        return None

    def build_deposition_exhibit_map(self, deponent_last_name: str) -> dict[str, str]:
        """
        Build a mapping of deposition exhibit numbers to Bates numbers
        for a specific deponent.

        Reads the "Deposition" column which contains entries like:
            "Smith Ex. 5; Jones Ex. 12"

        For deponent "Smith", this would produce: {"5": "ABC-00012345"}

        Args:
            deponent_last_name: Last name of the deponent (case-insensitive).

        Returns: Dict mapping exhibit number strings to Bates numbers.
                 e.g. {"5": "ABC-00012345", "12": "ABC-00067890"}
        """
        depo_col = self._col("Deposition")
        if not depo_col:
            logger.warning("No 'Deposition' column found in exhibit list.")
            return {}

        # Pattern: "LastName Ex. ##" — captures the exhibit number
        pattern = re.compile(
            rf"\b{re.escape(deponent_last_name)}\s+Ex\.\s*(\d+)\b",
            re.IGNORECASE,
        )

        exhibit_map = {}
        for row, bates in self.get_all_bates_numbers():
            depo_val = self._get_cell_value(row, "Deposition")
            if not depo_val:
                continue

            # Split by semicolon and check each entry
            for entry in depo_val.split(";"):
                match = pattern.search(entry.strip())
                if match:
                    exhibit_num = match.group(1)
                    exhibit_map[exhibit_num] = bates
                    logger.debug(f"  Mapped {deponent_last_name} Ex. {exhibit_num} → {bates}")

        logger.info(f"Built deposition exhibit map for '{deponent_last_name}': "
                     f"{len(exhibit_map)} exhibits.")
        return exhibit_map

    def build_all_deposition_exhibit_maps(self) -> dict[str, dict[str, str]]:
        """
        Build exhibit-to-Bates maps for ALL deponents found in the Deposition column.

        Returns: {deponent_last_name: {exhibit_num: bates_number, ...}, ...}
        """
        depo_col = self._col("Deposition")
        if not depo_col:
            return {}

        # First pass: collect all unique deponent names
        name_pattern = re.compile(r"(\w+)\s+Ex\.\s*\d+", re.IGNORECASE)
        deponent_names = set()

        for row, bates in self.get_all_bates_numbers():
            depo_val = self._get_cell_value(row, "Deposition")
            if not depo_val:
                continue
            for entry in depo_val.split(";"):
                match = name_pattern.search(entry.strip())
                if match:
                    deponent_names.add(match.group(1))

        # Second pass: build maps for each deponent
        all_maps = {}
        for name in deponent_names:
            all_maps[name.lower()] = self.build_deposition_exhibit_map(name)

        logger.info(f"Built exhibit maps for {len(all_maps)} deponents: "
                     f"{', '.join(sorted(all_maps.keys()))}")
        return all_maps

    def get_document_url(self, bates_number: str) -> Optional[str]:
        """Get the SharePoint URL for opening a document by its Bates number."""
        folder_files = self.client.list_files_in_folder(DOCUMENTS_FOLDER)
        if bates_number in folder_files:
            file_info = folder_files[bates_number]
            url = file_info["webUrl"]
            ext = file_info["extension"]
            if ext in (".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"):
                url = f"{url}?web=0"
            return url
        return None
